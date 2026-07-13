from pathlib import Path

import pytest

from app.chunking import build_chunks
from app.parsers.interface import ParsedBlock
from app.parsers.registry import parse_document
from app.parsers.sections import detect_section


def write_minimal_text_pdf(path: Path) -> None:
    content = (
        b"BT /F1 12 Tf 72 720 Td "
        b"(Abstract) Tj 0 -20 Td "
        b"(The sample was heated.) Tj ET"
    )
    objects = [
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n",
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n",
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>\nendobj\n",
        b"4 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n",
        b"5 0 obj\n<< /Length "
        + str(len(content)).encode()
        + b" >>\nstream\n"
        + content
        + b"\nendstream\nendobj\n",
    ]
    body = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for obj in objects:
        offsets.append(len(body))
        body.extend(obj)
    xref_at = len(body)
    body.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    body.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        body.extend(f"{offset:010d} 00000 n \n".encode())
    trailer = (
        f"trailer\n<< /Root 1 0 R /Size {len(objects) + 1} >>\n"
        f"startxref\n{xref_at}\n%%EOF\n"
    )
    body.extend(trailer.encode())
    path.write_bytes(body)


def test_markdown_parser_detects_headings_as_sections(tmp_path: Path):
    path = tmp_path / "paper.md"
    path.write_text(
        "# Abstract\nA short summary.\n\n"
        "## Materials & Methods\nWe mixed sodium chloride in water.\n\n"
        "## References\n[1] A paper.",
        encoding="utf-8",
    )

    parsed = parse_document(str(path), "text/markdown")

    assert parsed.errors == []
    assert parsed.pages[0].text.startswith("Abstract")
    assert [block.section for block in parsed.blocks if block.text and "sodium chloride" in block.text] == [
        "Materials and Methods"
    ]
    assert any(block.block_type == "heading" and block.section == "Abstract" for block in parsed.blocks)


def test_txt_parser_extracts_paragraph_blocks(tmp_path: Path):
    path = tmp_path / "notes.txt"
    path.write_text("Introduction\n\nFirst paragraph.\n\nSecond paragraph.", encoding="utf-8")

    parsed = parse_document(str(path), "text/plain")

    assert parsed.pages[0].page_number == 1
    assert [block.text for block in parsed.blocks if block.block_type == "paragraph"] == [
        "First paragraph.",
        "Second paragraph.",
    ]


def test_csv_parser_creates_table_block_and_text(tmp_path: Path):
    path = tmp_path / "assay.csv"
    path.write_text("compound,yield\nAspirin,75\nCaffeine,63\n", encoding="utf-8")

    parsed = parse_document(str(path), "text/csv")

    assert parsed.tables[0].csv_text == "compound,yield\nAspirin,75\nCaffeine,63\n"
    assert parsed.tables[0].rows == [
        {"compound": "Aspirin", "yield": 75},
        {"compound": "Caffeine", "yield": 63},
    ]
    assert parsed.blocks[0].block_type == "table"
    assert parsed.blocks[0].section == "Tables"
    assert parsed.blocks[0].metadata["row_count"] == 2
    assert "Aspirin" in parsed.pages[0].text


def test_xlsx_parser_creates_table_per_sheet(tmp_path: Path):
    from openpyxl import Workbook

    path = tmp_path / "workbook.xlsx"
    workbook = Workbook()
    assay = workbook.active
    assay.title = "Assay"
    assay.append(["compound", "yield"])
    assay.append(["A", 75])
    runs = workbook.create_sheet("Runs")
    runs.append(["sample", "temperature_c"])
    runs.append(["B", 80])
    workbook.save(path)

    parsed = parse_document(str(path), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    assert len(parsed.tables) == 2
    assert [table.metadata["sheet_name"] for table in parsed.tables] == ["Assay", "Runs"]
    assert [block.block_type for block in parsed.blocks] == ["table", "table"]
    assert parsed.blocks[1].metadata["rows"] == [{"sample": "B", "temperature_c": 80}]
    assert "Sheet: Runs" in parsed.pages[1].text


def test_pdf_parser_extracts_text_with_pypdf_fallback(tmp_path: Path):
    path = tmp_path / "paper.pdf"
    write_minimal_text_pdf(path)

    parsed = parse_document(str(path), "application/pdf")

    assert parsed.errors == []
    assert parsed.pages[0].page_number == 1
    assert "The sample was heated" in (parsed.pages[0].text or "")
    assert any(block.section == "Abstract" for block in parsed.blocks)


def test_section_detection_supports_common_scientific_variants():
    assert detect_section("Experimental Section") == "Experimental"
    assert detect_section("Materials & Methods") == "Materials and Methods"
    assert detect_section("SI") == "Supporting Information"
    assert detect_section("REFERENCES") == "References"


def test_chunking_groups_by_section_and_excludes_references():
    blocks = [
        ParsedBlock(block_type="paragraph", page_number=1, section="Experimental", text="alpha " * 20),
        ParsedBlock(block_type="paragraph", page_number=2, section="Experimental", text="beta " * 20),
        ParsedBlock(block_type="paragraph", page_number=3, section="References", text="[1] ref"),
        ParsedBlock(block_type="table", page_number=4, section=None, text="compound,yield\nA,75"),
    ]

    chunks = build_chunks(blocks, max_tokens=12)

    assert [chunk.chunk_index for chunk in chunks] == list(range(len(chunks)))
    assert {chunk.section for chunk in chunks} >= {"Experimental", "Tables"}
    assert "References" not in {chunk.section for chunk in chunks}
    assert all(chunk.token_count <= 12 for chunk in chunks)
    assert chunks[0].page_start == 1


def test_chunking_respects_budget_for_cjk_and_chemical_notation():
    blocks = [
        ParsedBlock(block_type="paragraph", page_number=1, section="Experimental", text="样品在氮气保护下加热至八十摄氏度"),
        ParsedBlock(block_type="paragraph", page_number=2, section="Experimental", text="CC(=O)OC1=CC=CC=C1C(=O)O Na2SO4"),
    ]

    chunks = build_chunks(blocks, max_tokens=8)

    assert len(chunks) >= 3
    assert all(0 < chunk.token_count <= 8 for chunk in chunks)
    assert "".join(chunk.text.replace(" ", "") for chunk in chunks).startswith("样品在氮气保护下")


def test_chunking_rejects_non_positive_budget():
    with pytest.raises(ValueError, match="max_tokens"):
        build_chunks([], max_tokens=0)
