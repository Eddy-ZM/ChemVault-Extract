from pathlib import Path
import csv
from datetime import date, datetime
from html import escape
import re

from app.parsers.interface import ParsedBlock, ParsedDocument, ParsedPage, ParsedTable


def _table_html(headers: list[str], rows: list[list[object]]) -> str:
    header_cells = "".join(f"<th>{escape(str(value))}</th>" for value in headers)
    body_rows = []
    for row in rows:
        body_rows.append("<tr>" + "".join(f"<td>{escape(str(value))}</td>" for value in row) + "</tr>")
    return (
        "<table><thead><tr>"
        + header_cells
        + "</tr></thead><tbody>"
        + "".join(body_rows)
        + "</tbody></table>"
    )


def _json_value(value: object) -> object:
    if value is None:
        return None
    try:
        import pandas as pd  # type: ignore

        if pd.isna(value):
            return None
    except Exception:
        pass
    if hasattr(value, "item"):
        try:
            value = value.item()
        except Exception:
            pass
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, str):
        stripped = value.strip()
        if stripped == "":
            return None
        if re.fullmatch(r"-?\d+", stripped):
            try:
                return int(stripped)
            except ValueError:
                return value
        if re.fullmatch(r"-?(\d+\.\d*|\d*\.\d+)", stripped):
            try:
                return float(stripped)
            except ValueError:
                return value
    if isinstance(value, int | float | str | bool):
        return value
    return str(value)


def _records_from_dataframe(df) -> list[dict]:
    records: list[dict] = []
    for row in df.to_dict(orient="records"):
        records.append({str(key): _json_value(value) for key, value in row.items()})
    return records


def _records_from_rows(headers: list[str], rows: list[list[object]]) -> list[dict]:
    records: list[dict] = []
    for row in rows:
        records.append(
            {
                str(header): _json_value(row[index] if index < len(row) else None)
                for index, header in enumerate(headers)
            }
        )
    return records


def _table_text(columns: list[str], rows: list[dict], csv_text: str, *, title: str | None = None) -> str:
    sample_rows = rows[:20]
    lines = []
    if title:
        lines.append(title)
    lines.append("Columns: " + ", ".join(str(column) for column in columns))
    if sample_rows:
        lines.append("Row samples:")
        for index, row in enumerate(sample_rows, start=1):
            values = "; ".join(f"{key}={value}" for key, value in row.items())
            lines.append(f"{index}. {values}")
    lines.append("CSV:")
    lines.append(csv_text.strip())
    return "\n".join(lines).strip()


def _normalize_csv_text(value: str) -> str:
    return value.replace("\r\n", "\n").replace("\r", "\n")


def parse_csv_file(file_path: str) -> ParsedDocument:
    path = Path(file_path)
    try:
        import pandas as pd  # type: ignore

        df = pd.read_csv(path)
        csv_text = _normalize_csv_text(df.to_csv(index=False))
        html = df.to_html(index=False)
        columns = [str(column) for column in df.columns]
        rows = _records_from_dataframe(df)
    except Exception:
        with path.open(newline="", encoding="utf-8", errors="replace") as file:
            raw_rows = list(csv.reader(file))
        headers = raw_rows[0] if raw_rows else []
        data_rows = raw_rows[1:] if len(raw_rows) > 1 else []
        csv_text = _normalize_csv_text(path.read_text(encoding="utf-8", errors="replace"))
        html = _table_html(headers, data_rows)
        columns = headers
        rows = _records_from_rows(headers, data_rows)

    metadata = {"parser": "csv", "filename": path.name, "row_count": len(rows), "columns": columns}
    text = _table_text(columns, rows, csv_text)
    table = ParsedTable(
        page_number=1,
        html=html,
        csv_text=csv_text,
        rows=rows,
        metadata=metadata,
        section="Tables",
    )
    block = ParsedBlock(
        block_type="table",
        text=text,
        page_number=1,
        section="Tables",
        html=html,
        metadata={**metadata, "rows": rows},
    )
    return ParsedDocument(
        metadata=metadata,
        pages=[ParsedPage(page_number=1, text=text, metadata={"source": "csv"})],
        blocks=[block],
        tables=[table],
    )


def parse_xlsx_file(file_path: str) -> ParsedDocument:
    path = Path(file_path)
    try:
        import pandas as pd  # type: ignore

        sheets = pd.read_excel(path, sheet_name=None)
        sheet_items = []
        for sheet_name, df in sheets.items():
            sheet_items.append(
                (
                    sheet_name,
                    _normalize_csv_text(df.to_csv(index=False)),
                    df.to_html(index=False),
                    [str(column) for column in df.columns],
                    _records_from_dataframe(df),
                )
            )
    except Exception:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet_items = []
        for worksheet in workbook.worksheets:
            values = list(worksheet.values)
            headers = [str(value or "") for value in values[0]] if values else []
            rows = [[value if value is not None else "" for value in row] for row in values[1:]]
            csv_lines = []
            if headers:
                csv_lines.append(",".join(headers))
            csv_lines.extend(",".join(str(value) for value in row) for row in rows)
            csv_text = _normalize_csv_text("\n".join(csv_lines) + ("\n" if csv_lines else ""))
            sheet_items.append(
                (
                    worksheet.title,
                    csv_text,
                    _table_html(headers, rows),
                    headers,
                    _records_from_rows(headers, rows),
                )
            )

    pages: list[ParsedPage] = []
    blocks: list[ParsedBlock] = []
    tables: list[ParsedTable] = []

    for index, (sheet_name, csv_text, html, columns, rows) in enumerate(sheet_items, start=1):
        metadata = {
            "parser": "xlsx",
            "filename": path.name,
            "sheet_name": sheet_name,
            "row_count": len(rows),
            "columns": columns,
        }
        text = _table_text(columns, rows, csv_text, title=f"Sheet: {sheet_name}")
        pages.append(ParsedPage(page_number=index, text=text, metadata={"sheet_name": sheet_name}))
        blocks.append(
            ParsedBlock(
                block_type="table",
                text=text,
                page_number=index,
                section="Tables",
                html=html,
                metadata={**metadata, "rows": rows},
            )
        )
        tables.append(
            ParsedTable(
                page_number=index,
                html=html,
                csv_text=csv_text,
                rows=rows,
                metadata=metadata,
                section="Tables",
            )
        )

    return ParsedDocument(
        metadata={"parser": "xlsx", "filename": path.name, "sheetCount": len(sheet_items)},
        pages=pages,
        blocks=blocks,
        tables=tables,
    )
